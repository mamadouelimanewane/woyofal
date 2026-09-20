"""
Génère les livrables du kit pilote dans docs/dist/ :
  - un PDF par document Markdown (GUIDE-UTILISATEUR, SCRIPT-DEMO, FICHE-PILOTE, DECLARATION-CDP)
  - les modèles d'import Excel (Immo et Entreprise)
Usage : python docs/build.py
Dépendances : fpdf2, markdown-it-py, openpyxl (déjà présents sur le poste).
"""
from __future__ import annotations
import re
from pathlib import Path
from fpdf import FPDF
from markdown_it import MarkdownIt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ICI = Path(__file__).parent
DIST = ICI / "dist"
DIST.mkdir(exist_ok=True)
FONTS = Path("C:/Windows/Fonts")
TEAL = (15, 118, 110)
GRIS = (100, 110, 120)
NOIR = (25, 30, 38)


class Doc(FPDF):
    def __init__(self, titre: str):
        super().__init__(format="A4")
        self.titre = titre
        self.set_margins(18, 18, 18)
        self.set_auto_page_break(True, 20)
        self.add_font("UI", "", str(FONTS / "segoeui.ttf"))
        self.add_font("UI", "B", str(FONTS / "segoeuib.ttf"))
        self.add_font("UI", "I", str(FONTS / "segoeuii.ttf"))
        self.add_font("Mono", "", str(FONTS / "consola.ttf"))

    def header(self):
        self.set_font("UI", "B", 9)
        self.set_text_color(*TEAL)
        self.cell(0, 6, "KURAÑ", new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(*TEAL)
        self.line(18, 24, 192, 24)
        self.ln(4)

    def footer(self):
        self.set_y(-14)
        self.set_font("UI", "", 8)
        self.set_text_color(*GRIS)
        self.cell(0, 6, f"{self.titre}  ·  processingenierie  ·  page {self.page_no()}", align="C")


def inline_html(text: str) -> str:
    """Markdown inline minimal → HTML compris par fpdf2 (gras, italique, code, liens)."""
    text = re.sub(r"`([^`]+)`", r"<font face='Mono'>\1</font>", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"<a href='\2'>\1</a>", text)
    return text


def plain(text: str) -> str:
    return re.sub(r"\*\*|`|(?<!\*)\*(?!\*)", "", re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text))


def tableau(pdf: Doc, lignes: list[list[str]]):
    if not lignes:
        return
    ncol = len(lignes[0])
    largeur = pdf.w - 36
    # largeur proportionnelle à la longueur moyenne des cellules, bornée
    poids = [max(8, min(60, sum(len(plain(l[i])) for l in lignes if i < len(l)) / len(lignes))) for i in range(ncol)]
    total = sum(poids)
    larg = [largeur * p / total for p in poids]
    pdf.set_font("UI", "", 8.5)
    for idx, ligne in enumerate(lignes):
        cellules = [plain(ligne[i]) if i < len(ligne) else "" for i in range(ncol)]
        # hauteur = max des lignes de texte
        h_ligne = 4.6
        nb = [max(1, len(pdf.multi_cell(larg[i], h_ligne, cellules[i], dry_run=True, output="LINES"))) for i in range(ncol)]
        h = max(nb) * h_ligne + 2
        if pdf.get_y() + h > pdf.h - 20:
            pdf.add_page()
        x0, y0 = pdf.get_x(), pdf.get_y()
        if idx == 0:
            pdf.set_fill_color(230, 244, 242)
            pdf.set_font("UI", "B", 8.5)
        elif idx % 2 == 0:
            pdf.set_fill_color(248, 250, 250)
        else:
            pdf.set_fill_color(255, 255, 255)
        for i in range(ncol):
            pdf.set_xy(x0 + sum(larg[:i]), y0)
            pdf.set_text_color(*NOIR)
            pdf.multi_cell(larg[i], h_ligne, cellules[i], border=0, fill=True, max_line_height=h_ligne, new_x="RIGHT", new_y="TOP")
            # remplissage de la hauteur de ligne
            pdf.set_xy(x0 + sum(larg[:i]), y0)
            pdf.cell(larg[i], h, "", border="B", fill=False)
        pdf.set_font("UI", "", 8.5)
        pdf.set_xy(x0, y0 + h)
    pdf.ln(3)


def md_vers_pdf(source: Path, titre: str):
    pdf = Doc(titre)
    pdf.add_page()
    lignes = source.read_text(encoding="utf-8").splitlines()
    i = 0
    premier_titre = True
    while i < len(lignes):
        l = lignes[i]
        if l.startswith("|") and i + 1 < len(lignes) and re.match(r"^\|[\s:-]+\|", lignes[i + 1]):
            bloc = []
            while i < len(lignes) and lignes[i].startswith("|"):
                if not re.match(r"^\|[\s:-]+\|", lignes[i]):
                    bloc.append([c.strip() for c in lignes[i].strip().strip("|").split("|")])
                i += 1
            tableau(pdf, bloc)
            continue
        if l.startswith("# "):
            pdf.set_font("UI", "B", 20 if premier_titre else 16)
            pdf.set_text_color(*TEAL)
            pdf.multi_cell(0, 9, plain(l[2:]), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
            premier_titre = False
        elif l.startswith("## "):
            if pdf.get_y() > pdf.h - 50:
                pdf.add_page()
            pdf.ln(3)
            pdf.set_font("UI", "B", 13)
            pdf.set_text_color(*TEAL)
            pdf.multi_cell(0, 7, plain(l[3:]), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif l.startswith("### "):
            pdf.ln(2)
            pdf.set_font("UI", "B", 11)
            pdf.set_text_color(*NOIR)
            pdf.multi_cell(0, 6, plain(l[4:]), new_x="LMARGIN", new_y="NEXT")
        elif l.startswith("---"):
            pdf.ln(2)
            pdf.set_draw_color(210, 215, 220)
            pdf.line(18, pdf.get_y(), 192, pdf.get_y())
            pdf.ln(3)
        elif l.startswith("> "):
            pdf.set_font("UI", "I", 9.5)
            pdf.set_text_color(*GRIS)
            pdf.set_x(24)
            pdf.multi_cell(0, 5.2, plain(l[2:]), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif re.match(r"^\s*[-*] ", l) or re.match(r"^\s*\d+\. ", l) or l.startswith("- [ ]"):
            texte = re.sub(r"^\s*(?:[-*]|\d+\.)\s+", "", l)
            puce = "☐ " if texte.startswith("[ ] ") else ("• " if re.match(r"^\s*[-*] ", l) else re.match(r"^\s*(\d+)\.", l).group(1) + ". ")
            texte = texte.replace("[ ] ", "")
            pdf.set_font("UI", "", 10)
            pdf.set_text_color(*NOIR)
            pdf.set_x(22)
            pdf.write_html(f"<font size='10'>{puce}{inline_html(texte)}</font>")
            pdf.ln(5.2)
        elif l.strip() == "":
            pdf.ln(1.5)
        else:
            pdf.set_font("UI", "", 10)
            pdf.set_text_color(*NOIR)
            pdf.write_html(f"<font size='10'>{inline_html(l)}</font>")
            pdf.ln(5.4)
        i += 1
    cible = DIST / (source.stem + ".pdf")
    pdf.output(str(cible))
    print("PDF  ->", cible.name)


def modele_excel(nom: str, immo: bool):
    wb = Workbook()
    entete = Font(bold=True, color="FFFFFF")
    fond = PatternFill("solid", fgColor="0F766E")
    onglets = {
        "Sites": (["nom", "type", "adresse", "surfaceM2", "budgetMensuel"],
                  [["Immeuble Sérigne Fallou", "immeuble", "Sacré-Cœur 3, Dakar", 620, ""], ["Cour commune Médina", "cour commune", "Rue 22 x 31, Médina", 280, ""]] if immo
                  else [["Pharmacie Liberté 6", "pharmacie", "Liberté 6, Dakar", 72, 110000], ["Pharmacie Pikine", "pharmacie", "Pikine, route de Thiaroye", 60, 90000]]),
        "Lots": (["site", "reference", "surfaceM2", "etage"],
                 [["Immeuble Sérigne Fallou", "Appt 1", 85, "1"], ["Immeuble Sérigne Fallou", "Appt 2", 110, "1"], ["Cour commune Médina", "Chambre 1", 20, ""], ["Cour commune Médina", "Chambre 2", 20, ""], ["Cour commune Médina", "Boutique", 30, ""]] if immo
                 else [["Pharmacie Liberté 6", "Officine", "", ""], ["Pharmacie Pikine", "Officine", "", ""]]),
        "Compteurs": (["site", "numero", "typeTarif", "libelle", "lots", "regle"],
                      [["Immeuble Sérigne Fallou", "14210034561", "DPP", "Appt 1", "Appt 1", "egal"], ["Immeuble Sérigne Fallou", "14210034562", "DPP", "Appt 2", "Appt 2", "egal"], ["Immeuble Sérigne Fallou", "14210034599", "DPP", "Parties communes", "Appt 1; Appt 2", "surface"], ["Cour commune Médina", "14210034567", "DPP", "Compteur unique", "Chambre 1; Chambre 2; Boutique", "egal"]] if immo
                      else [["Pharmacie Liberté 6", "14200011122", "PRO", "Officine", "Officine", ""], ["Pharmacie Pikine", "14200011123", "PRO", "Officine", "Officine", ""]]),
        "Occupants": (["site", "lot", "nom", "telephone", "dateEntree", "caution"],
                      [["Immeuble Sérigne Fallou", "Appt 1", "Mariama Diallo", "77 201 00 01", "2025-02-01", 25000], ["Cour commune Médina", "Chambre 1", "Khady Thiam", "76 301 11 11", "2025-03-15", 10000]] if immo else []),
    }
    wb.remove(wb.active)
    for onglet, (cols, exemples) in onglets.items():
        ws = wb.create_sheet(onglet)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = entete, fond
        for ex in exemples:
            ws.append(ex)
        for i, col in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(14, len(col) + 6)
        ws.freeze_panes = "A2"
    aide = wb.create_sheet("Aide")
    for l in [
        ["KURAÑ — modèle d'import du parc"],
        ["Remplissez les onglets Sites, Lots, Compteurs, Occupants (les exemples peuvent être effacés)."],
        ["numero : 11 chiffres, tel qu'inscrit sur le compteur et sur les tickets."],
        ["typeTarif : DPP (domestique petite puissance), DMP (moyenne puissance), PRO (professionnel)."],
        ["lots : références de lots séparées par « ; » — plusieurs lots = compteur partagé."],
        ["regle : egal, surface, sous_compteur ou forfait (vide = parts égales)."],
        ["dateEntree : AAAA-MM-JJ. telephone : 9 chiffres (77 000 00 00)."],
        ["Les sites et lots déjà présents dans KURAÑ sont reconnus par leur nom et ne sont jamais dupliqués."],
    ]:
        aide.append(l)
    aide.column_dimensions["A"].width = 110
    aide["A1"].font = Font(bold=True, size=13, color="0F766E")
    cible = DIST / nom
    wb.save(cible)
    print("XLSX ->", cible.name)


if __name__ == "__main__":
    md_vers_pdf(ICI / "GUIDE-UTILISATEUR.md", "Guide utilisateur")
    md_vers_pdf(ICI / "SCRIPT-DEMO.md", "Script de démonstration")
    md_vers_pdf(ICI / "FICHE-PILOTE.md", "Programme pilote")
    md_vers_pdf(ICI / "DECLARATION-CDP.md", "Déclaration CDP")
    modele_excel("kuran-modele-import-immo.xlsx", True)
    modele_excel("kuran-modele-import-entreprise.xlsx", False)
